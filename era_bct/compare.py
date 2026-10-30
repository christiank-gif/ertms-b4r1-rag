from __future__ import annotations

import math


# SECTION 11/13 — VALIDATION CELLULE PAR CELLULE CONTRE LE CLASSEUR (valeurs en cache du VBA).
# ==========================================================================

"""Cell-by-cell comparison of a Python run with the results cached in a workbook
(the values the VBA macro wrote the last time 'Calculate the curves' was pressed)."""

import math


def _eq(a, b, tol):
    if a in (None, "") and b in (None, ""):
        return True
    if isinstance(a, bool) or isinstance(b, bool):
        return bool(a) == bool(b)
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return math.isclose(a, b, rel_tol=tol, abs_tol=tol)
    return a == b


# Validation cellule par cellule contre les valeurs en cache du classeur (celles
# écrites par le VBA au dernier 'Calculate the curves').
def compare_with_workbook(result, path, tol=1e-9, max_report=20) -> dict:
    """Compares Curves (N2:AZ), Track (A36:F), deceleration table and Calc pivots."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    e = result.engine
    report = {}

    def cmp(name, ws_name, grid_get, rows, cols):
        ws = wb[ws_name]
        n = bad = 0
        examples = []
        for r_idx, row in enumerate(ws.iter_rows(min_row=rows[0], max_row=rows[1], min_col=cols[0],
                                                 max_col=cols[1], values_only=True), start=rows[0]):
            for c_idx, xv in enumerate(row, start=cols[0]):
                pv = grid_get(r_idx, c_idx)
                if xv in (None, "") and pv in (None, ""):
                    continue
                n += 1
                if not _eq(pv, xv, tol):
                    bad += 1
                    if len(examples) < max_report:
                        examples.append((r_idx, c_idx, xv, pv))
        report[name] = {"cells": n, "mismatches": bad, "examples": examples}

    sheet = "Curves Gamma train" if e.gamma else "Curves Lambda train"
    last = max([max(c) for c in e.curves.cols.values() if c] + [2])
    cmp("curves", sheet, lambda r, c: e.curves[r, c], (2, last + 5), (14, 52))
    tlast = max(e.track.cols.get(1, {0: 0}))
    cmp("track", "Track", lambda r, c: e.track[r, c], (36, tlast + 5), (1, 6))
    dsheet = "Gamma train deceleration" if e.gamma else "Lambda train deceleration"
    dcols = 19 if e.gamma else 10
    cmp("deceleration", dsheet, lambda r, c: e.decel[r, c], (2, max(e.decel.cols[1]) + 1), (1, dcols))
    csheet = "Calc Gamma" if e.gamma else "Calc Lambda"
    pcols = (1, 10) if e.gamma else (1, 7)
    cmp("pivots", csheet, lambda r, c: e.calc[r, c], (2, 44), pcols)
    ws = wb[sheet]
    xl = [ws.cell(38, c).value for c in range(1, 10)]
    s = result.summary
    py = [s.get(k) for k in ("V_linespeed", "perturbation", "indication", "permitted", "warning", "sbi", "ebi",
                             "start_rsm", "release_speed")]
    report["summary"] = {"excel": xl, "python": py, "match": all(_eq(a, b, tol) for a, b in zip(py, xl))}
    return report


# ==========================================================================
